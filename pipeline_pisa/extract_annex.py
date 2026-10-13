import csv
import hashlib
import json
from pathlib import Path
import re
import sys
import openpyxl

SOURCE_URL = "https://stat.link/files/73451bc5-en/mrq53f.xlsx"
SOURCE_LICENCE = "CC BY 4.0"
OUTPUT_COLUMNS = ["table", "subject", "domain", "stat", "cycle", "value", "se"]


def compute_sha256(filepath):
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def find_means_columns(ws, sheet_name, assertions):
    cols = []
    # Mean score
    mean_col = None
    for c in range(1, ws.max_column + 1):
        r7 = str(ws.cell(7, c).value or "").strip().lower()
        r9 = str(ws.cell(9, c).value or "").strip().lower()
        if "mean score" in r7 and r9 == "mean":
            se_val = str(ws.cell(9, c + 1).value or "").strip()
            if se_val != "S.E.":
                sys.exit(f"Assertion failed: {sheet_name} col {c+1} expected S.E., got {se_val}")
            mean_col = c
            cols.append(("mean", c, c + 1))
            break
    if mean_col is None:
        sys.exit(f"Assertion failed: {sheet_name} missing 'Mean score' header")
    assertions.append(f"{sheet_name}: found header 'Mean score' at column {mean_col}")

    # Standard deviation
    sd_col = None
    for c in range(1, ws.max_column + 1):
        r7 = str(ws.cell(7, c).value or "").strip().lower()
        r9 = str(ws.cell(9, c).value or "").strip().lower()
        if "standard deviation" in r7 and r9 == "s.d.":
            se_val = str(ws.cell(9, c + 1).value or "").strip()
            if se_val != "S.E.":
                sys.exit(f"Assertion failed: {sheet_name} col {c+1} expected S.E., got {se_val}")
            sd_col = c
            cols.append(("sd", c, c + 1))
            break
    if sd_col is None:
        sys.exit(f"Assertion failed: {sheet_name} missing 'Standard deviation' header")
    assertions.append(f"{sheet_name}: found header 'Standard deviation' at column {sd_col}")

    # Percentiles
    pct_map = [
        ("10th", "p10"),
        ("25th", "p25"),
        ("50th", "p50"),
        ("75th", "p75"),
        ("90th", "p90"),
    ]
    for label_sub, stat_name in pct_map:
        found_col = None
        for c in range(1, ws.max_column + 1):
            r8 = str(ws.cell(8, c).value or "").strip()
            if "diff" in r8.lower():
                continue
            if label_sub in r8:
                se_val = str(ws.cell(9, c + 1).value or "").strip()
                if se_val != "S.E.":
                    sys.exit(f"Assertion failed: {sheet_name} percentile {stat_name} missing S.E. at col {c+1}")
                found_col = c
                cols.append((stat_name, c, c + 1))
                break
        if found_col is None:
            sys.exit(f"Assertion failed: {sheet_name} missing percentile '{stat_name}' header")
        assertions.append(f"{sheet_name}: found percentile '{stat_name}' ({label_sub}) at column {found_col}")

    return cols


def find_proficiency_columns(ws, sheet_name, domain, assertions):
    bands = []
    for c in range(1, ws.max_column):
        r8 = ws.cell(8, c).value
        if r8 and ("level" in str(r8).lower()):
            first_line = str(r8).split("\n")[0].strip()
            m = re.match(r"^(Below\s+Level|Level)\s+([0-9a-c]+)$", first_line, re.IGNORECASE)
            if m:
                prefix, num = m.groups()
                stat_token = f"level_below_{num.lower()}" if "below" in prefix.lower() else f"level_{num.lower()}"
                v_label = str(ws.cell(9, c).value or "").strip()
                se_label = str(ws.cell(9, c + 1).value or "").strip()
                if v_label != "%":
                    sys.exit(f"Assertion failed: {sheet_name} col {c} expected '%', got '{v_label}'")
                if se_label != "S.E.":
                    sys.exit(f"Assertion failed: {sheet_name} col {c+1} expected 'S.E.', got '{se_label}'")
                bands.append((stat_token, c, c + 1, first_line))

    expected_count = 8 if domain == "science" else 9
    if len(bands) != expected_count:
        sys.exit(f"Assertion failed: {sheet_name} expected {expected_count} bands, found {len(bands)}")

    assertions.append(f"{sheet_name}: found {len(bands)} proficiency bands ({', '.join(b[0] for b in bands)})")
    return bands


def find_lowtop_columns(ws, sheet_name, assertions):
    cols = []
    for c in range(1, ws.max_column + 1):
        r7 = ws.cell(7, c).value
        if r7 and "proficiency level in pisa" in str(r7).lower():
            m = re.search(r"pisa\s+(\d{4})", str(r7), re.IGNORECASE)
            if not m:
                sys.exit(f"Assertion failed: {sheet_name} could not parse cycle from '{r7}'")
            cycle = int(m.group(1))

            # Column c: below level 2
            r8_below = str(ws.cell(8, c).value or "").strip()
            if "below level 2" not in r8_below.lower():
                sys.exit(f"Assertion failed: {sheet_name} col {c} expected 'Below Level 2', got '{r8_below}'")
            if ws.cell(9, c).value != "%" or ws.cell(9, c + 1).value != "S.E.":
                sys.exit(f"Assertion failed: {sheet_name} col {c} expected % and S.E. headers")
            cols.append(("below_l2", cycle, c, c + 1))

            # Column c+2: level 5 or above
            r8_above = str(ws.cell(8, c + 2).value or "").strip()
            if "level 5 or above" not in r8_above.lower():
                sys.exit(f"Assertion failed: {sheet_name} col {c+2} expected 'Level 5 or above', got '{r8_above}'")
            if ws.cell(9, c + 2).value != "%" or ws.cell(9, c + 3).value != "S.E.":
                sys.exit(f"Assertion failed: {sheet_name} col {c+2} expected % and S.E. headers")
            cols.append(("above_l4", cycle, c + 2, c + 3))

    if not cols:
        sys.exit(f"Assertion failed: {sheet_name} no lowtop cycle columns found")
    cycles = sorted(list({col[1] for col in cols}))
    assertions.append(f"{sheet_name}: found lowtop columns for cycles {cycles}")
    return cols


def find_trend_science_columns(ws, sheet_name, assertions):
    cols = []
    r7 = str(ws.cell(7, 2).value or "").strip()
    if "science performance, by pisa cycle" not in r7.lower():
        sys.exit(f"Assertion failed: {sheet_name} missing 'Science performance, by PISA cycle' header")

    for c in range(1, ws.max_column + 1):
        r8 = ws.cell(8, c).value
        if r8 and "pisa" in str(r8).lower() and "-" not in str(r8) and "diff" not in str(r8).lower():
            if "change" in str(ws.cell(7, c).value or "").lower():
                continue
            m = re.match(r"^\s*pisa\s+(\d{4})\s*$", str(r8), re.IGNORECASE)
            if m:
                cycle = int(m.group(1))
                r9_v = str(ws.cell(9, c).value or "").strip().lower()
                r9_se = str(ws.cell(9, c + 1).value or "").strip().lower()
                if "mean score" not in r9_v:
                    sys.exit(f"Assertion failed: {sheet_name} col {c} expected 'Mean score', got '{r9_v}'")
                if "s.e." not in r9_se:
                    sys.exit(f"Assertion failed: {sheet_name} col {c+1} expected 'S.E.', got '{r9_se}'")
                cols.append(("mean", cycle, c, c + 1))

    if not cols:
        sys.exit(f"Assertion failed: {sheet_name} no trend science columns found")
    cycles = sorted(list({col[1] for col in cols}))
    assertions.append(f"{sheet_name}: found trend science columns for cycles {cycles}")
    return cols


def find_variation_science_columns(ws, sheet_name, assertions):
    cols = []
    # Total variance
    tot_col = None
    for c in range(1, ws.max_column + 1):
        r6 = str(ws.cell(6, c).value or "").strip().lower()
        if "total variation in science performance" in r6:
            if ws.cell(9, c).value != "Variance" or ws.cell(9, c + 1).value != "S.E.":
                sys.exit(f"Assertion failed: {sheet_name} total_variance col {c} missing Variance/S.E.")
            tot_col = c
            cols.append(("total_variance", c, c + 1))
            break
    if tot_col is None:
        sys.exit(f"Assertion failed: {sheet_name} missing 'Total variation in science performance' header")
    assertions.append(f"{sheet_name}: found 'total_variance' at column {tot_col}")

    # Between school variance
    btw_col = None
    for c in range(1, ws.max_column + 1):
        r6 = str(ws.cell(6, c).value or "").strip().lower()
        if "between schools" in r6:
            if ws.cell(9, c).value != "Variance" or ws.cell(9, c + 1).value != "S.E.":
                sys.exit(f"Assertion failed: {sheet_name} between_school_variance col {c} missing Variance/S.E.")
            btw_col = c
            cols.append(("between_school_variance", c, c + 1))
            break
    if btw_col is None:
        sys.exit(f"Assertion failed: {sheet_name} missing 'Variation in science performance between schools' header")
    assertions.append(f"{sheet_name}: found 'between_school_variance' at column {btw_col}")

    # Within school variance
    wth_col = None
    for c in range(1, ws.max_column + 1):
        r6 = str(ws.cell(6, c).value or "").strip().lower()
        if "within schools" in r6:
            if ws.cell(9, c).value != "Variance" or ws.cell(9, c + 1).value != "S.E.":
                sys.exit(f"Assertion failed: {sheet_name} within_school_variance col {c} missing Variance/S.E.")
            wth_col = c
            cols.append(("within_school_variance", c, c + 1))
            break
    if wth_col is None:
        sys.exit(f"Assertion failed: {sheet_name} missing 'Variation in science performance within schools' header")
    assertions.append(f"{sheet_name}: found 'within_school_variance' at column {wth_col}")

    # Between school pct of oecd average
    shr_col = None
    for c in range(1, ws.max_column + 1):
        r8 = str(ws.cell(8, c).value or "").strip().lower()
        if "between-school variation" in r8:
            if ws.cell(9, c).value != "%":
                sys.exit(f"Assertion failed: {sheet_name} between_school_pct_of_oecd_average col {c} expected '%'")
            shr_col = c
            cols.append(("between_school_pct_of_oecd_average", c, None))
            break
    if shr_col is None:
        sys.exit(f"Assertion failed: {sheet_name} missing 'Between-school variation' header")
    assertions.append(f"{sheet_name}: found 'between_school_pct_of_oecd_average' at column {shr_col}")

    return cols


def extract_subject_rows(ws, sheet_name, is_trend):
    subjects = []
    for r in range(10, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if not v or not str(v).strip():
            continue
        v_str = str(v).strip()
        if v_str.startswith(("Information on data", "Note:", "Notes:", "1.", "2.", "3.", "4.")):
            continue
        if any(kw in v_str for kw in ("conducted the PISA", "results reported under")):
            continue
        has_val = any(ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip() != "" for c in range(2, ws.max_column + 1))
        if has_val:
            subjects.append((r, v_str))

    sub_names = {sub[1] for sub in subjects}
    if "OECD average" not in sub_names:
        sys.exit(f"Assertion failed: {sheet_name} missing 'OECD average'")
    if "Indonesia" not in sub_names:
        sys.exit(f"Assertion failed: {sheet_name} missing 'Indonesia'")
    if is_trend and "OECD average-23" not in sub_names:
        sys.exit(f"Assertion failed: {sheet_name} missing 'OECD average-23'")

    return subjects


def format_value_and_se(val, se):
    if val in ("m", "c"):
        return None, None, val
    if not isinstance(val, (int, float)):
        return None, None, "empty"
    val_str = repr(float(val))

    if se is None or se in ("", "m", "c"):
        se_str = ""
    elif isinstance(se, (int, float)):
        se_str = repr(float(se))
    else:
        se_str = ""

    return val_str, se_str, None


def main():
    input_file = sys.argv[1] if len(sys.argv) > 1 else "/root/data/pisa2025/mrq53f.xlsx"
    input_path = Path(input_file)
    if not input_path.exists():
        sys.exit(f"Input file not found: {input_path}")

    input_sha256 = compute_sha256(input_path)

    wb = openpyxl.load_workbook(input_path, data_only=True)
    required_sheets = [
        "Table I.B1.2a.1", "Table I.B1.2a.2", "Table I.B1.2a.3",
        "Table I.B1.2a.10", "Table I.B1.2a.11", "Table I.B1.2a.12",
        "Table I.B1.2a.33", "Table I.B1.2a.34", "Table I.B1.2a.35",
        "Table I.B1.2a.36",
        "Table I.B1.2a.25",
    ]

    for sname in required_sheets:
        if sname not in wb.sheetnames:
            sys.exit(f"Assertion failed: sheet '{sname}' does not exist in workbook")

    assertions = []
    output_rows = {
        "annex_means.csv": [],
        "annex_proficiency.csv": [],
        "annex_trend_science.csv": [],
        "annex_lowtop.csv": [],
        "annex_variation_science.csv": [],
    }
    sheet_counts = {}
    sheet_skipped = {}

    # 1. Means sheets
    means_specs = [
        ("Table I.B1.2a.1", "science"),
        ("Table I.B1.2a.2", "reading"),
        ("Table I.B1.2a.3", "mathematics"),
    ]
    for sname, domain in means_specs:
        ws = wb[sname]
        cols = find_means_columns(ws, sname, assertions)
        subjects = extract_subject_rows(ws, sname, is_trend=False)
        assertions.append(f"{sname}: verified {len(subjects)} subjects found including Indonesia")

        written = 0
        skip_m = 0
        skip_c = 0
        for r, subject in subjects:
            for stat, vc, sec in cols:
                val = ws.cell(r, vc).value
                se = ws.cell(r, sec).value
                val_str, se_str, sym = format_value_and_se(val, se)
                if sym == "m":
                    skip_m += 1
                elif sym == "c":
                    skip_c += 1
                elif val_str is not None:
                    output_rows["annex_means.csv"].append([sname, subject, domain, stat, "latest", val_str, se_str])
                    written += 1
        sheet_counts[sname] = written
        sheet_skipped[sname] = {"m": skip_m, "c": skip_c}

    # 2. Proficiency sheets
    prof_specs = [
        ("Table I.B1.2a.10", "science"),
        ("Table I.B1.2a.11", "reading"),
        ("Table I.B1.2a.12", "mathematics"),
    ]
    for sname, domain in prof_specs:
        ws = wb[sname]
        bands = find_proficiency_columns(ws, sname, domain, assertions)
        subjects = extract_subject_rows(ws, sname, is_trend=False)
        assertions.append(f"{sname}: verified {len(subjects)} subjects found including Indonesia")

        written = 0
        skip_m = 0
        skip_c = 0
        for r, subject in subjects:
            for stat, vc, sec, _ in bands:
                val = ws.cell(r, vc).value
                se = ws.cell(r, sec).value
                val_str, se_str, sym = format_value_and_se(val, se)
                if sym == "m":
                    skip_m += 1
                elif sym == "c":
                    skip_c += 1
                elif val_str is not None:
                    output_rows["annex_proficiency.csv"].append([sname, subject, domain, stat, "latest", val_str, se_str])
                    written += 1
        sheet_counts[sname] = written
        sheet_skipped[sname] = {"m": skip_m, "c": skip_c}

    # 3. Lowtop sheets
    lowtop_specs = [
        ("Table I.B1.2a.33", "science"),
        ("Table I.B1.2a.34", "reading"),
        ("Table I.B1.2a.35", "mathematics"),
    ]
    for sname, domain in lowtop_specs:
        ws = wb[sname]
        cols = find_lowtop_columns(ws, sname, assertions)
        subjects = extract_subject_rows(ws, sname, is_trend=True)
        assertions.append(f"{sname}: verified {len(subjects)} subjects found including Indonesia and OECD average-23")

        written = 0
        skip_m = 0
        skip_c = 0
        for r, subject in subjects:
            for stat, cycle, vc, sec in cols:
                val = ws.cell(r, vc).value
                se = ws.cell(r, sec).value
                val_str, se_str, sym = format_value_and_se(val, se)
                if sym == "m":
                    skip_m += 1
                elif sym == "c":
                    skip_c += 1
                elif val_str is not None:
                    output_rows["annex_lowtop.csv"].append([sname, subject, domain, stat, cycle, val_str, se_str])
                    written += 1
        sheet_counts[sname] = written
        sheet_skipped[sname] = {"m": skip_m, "c": skip_c}

    # 4. Trend science sheet
    sname = "Table I.B1.2a.36"
    ws = wb[sname]
    cols = find_trend_science_columns(ws, sname, assertions)
    subjects = extract_subject_rows(ws, sname, is_trend=True)
    assertions.append(f"{sname}: verified {len(subjects)} subjects found including Indonesia and OECD average-23")

    written = 0
    skip_m = 0
    skip_c = 0
    for r, subject in subjects:
        for stat, cycle, vc, sec in cols:
            val = ws.cell(r, vc).value
            se = ws.cell(r, sec).value
            val_str, se_str, sym = format_value_and_se(val, se)
            if sym == "m":
                skip_m += 1
            elif sym == "c":
                skip_c += 1
            elif val_str is not None:
                output_rows["annex_trend_science.csv"].append([sname, subject, "science", stat, cycle, val_str, se_str])
                written += 1
    sheet_counts[sname] = written
    sheet_skipped[sname] = {"m": skip_m, "c": skip_c}

    # 5. Variation science sheet
    sname = "Table I.B1.2a.25"
    ws = wb[sname]
    cols = find_variation_science_columns(ws, sname, assertions)
    subjects = extract_subject_rows(ws, sname, is_trend=False)
    assertions.append(f"{sname}: verified {len(subjects)} subjects found including Indonesia")

    written = 0
    skip_m = 0
    skip_c = 0
    for r, subject in subjects:
        for stat, vc, sec in cols:
            val = ws.cell(r, vc).value
            se = ws.cell(r, sec).value if sec is not None else None
            val_str, se_str, sym = format_value_and_se(val, se)
            if sym == "m":
                skip_m += 1
            elif sym == "c":
                skip_c += 1
            elif val_str is not None:
                output_rows["annex_variation_science.csv"].append([sname, subject, "science", stat, "latest", val_str, se_str])
                written += 1
    sheet_counts[sname] = written
    sheet_skipped[sname] = {"m": skip_m, "c": skip_c}

    # All assertions passed, write CSV files
    out_dir = Path("/root/projects/dashboard/data/pisa")
    out_dir.mkdir(parents=True, exist_ok=True)

    for csv_filename, rows in output_rows.items():
        csv_path = out_dir / csv_filename
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f, lineterminator="\n")
            writer.writerow(OUTPUT_COLUMNS)
            for row in rows:
                writer.writerow(row)

    # Write report json
    report_data = {
        "input_sha256": input_sha256,
        "label_assertions_passed": assertions,
        "provenance": {
            "licence": SOURCE_LICENCE,
            "sha256": input_sha256,
            "source_url": SOURCE_URL,
        },
        "row_counts": sheet_counts,
        "skipped_counts": sheet_skipped,
    }
    report_path = out_dir / "annex_report.json"
    with open(report_path, "w", encoding="utf-8") as f:
        json.dump(report_data, f, indent=2, sort_keys=True)
        f.write("\n")

    # Print one line per sheet with tidy row count
    for sheet_name in required_sheets:
        print(f"{sheet_name}: {sheet_counts[sheet_name]} tidy rows written")


if __name__ == "__main__":
    main()
